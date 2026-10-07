"""Reportes tabulares y exportacion a CSV, Excel y PDF. Cada reporte declara sus columnas y devuelve filas."""
import csv
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO, StringIO
from typing import Callable

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.parametros import parametro
from app.core.util import ahora, dec, dia_local, hoy, inicio_dia_utc, json_valor
from app.modelos import (Bodega, ConteoInventario, DetalleConteoInventario, DocumentoInventario, ExcepcionFifo, Existencia,
                         Familia, Lote, MovimientoInventario, Prestamo, Proceso, Producto, Proveedor, ReservaStock,
                         Subfamilia, TipoMovimiento, Ubicacion, Usuario)

from . import costeo


@dataclass
class Filtros:
    empresa_id: int
    desde: date | None = None
    hasta: date | None = None
    bodega_id: int | None = None
    familia_id: int | None = None
    subfamilia_id: int | None = None
    producto_id: int | None = None
    proceso_id: int | None = None
    temporada_id: int | None = None
    especie: str | None = None
    productor: str | None = None
    proveedor_id: int | None = None
    dias: int | None = None
    bodegas: set[int] | None = None  # alcance del usuario
    familias: set[int] | None = None


@dataclass
class Reporte:
    codigo: str
    titulo: str
    columnas: list[tuple[str, str, str]]  # (clave, titulo, tipo: texto|numero|moneda|fecha|fechahora)
    generar: Callable[[Session, Filtros], list[dict]]
    requiere_costos: bool = False
    descripcion: str = ""
    filtros: list[str] = field(default_factory=list)


# ---------------------------------------------------------------- consultas base

def existencias_con_stock(db: Session, f: Filtros, solo_utilizables: bool = False) -> list[Existencia]:
    c = (select(Existencia).join(Producto, Existencia.producto_id == Producto.id)
         .join(Subfamilia, Producto.subfamilia_id == Subfamilia.id)
         .where(Existencia.empresa_id == f.empresa_id, Existencia.cantidad_fisica > 0, Existencia.estado != "ANULADA"))
    for col, val in ((Existencia.bodega_id, f.bodega_id), (Subfamilia.familia_id, f.familia_id),
                     (Producto.subfamilia_id, f.subfamilia_id), (Existencia.producto_id, f.producto_id),
                     (Existencia.proveedor_id, f.proveedor_id)):
        if val:
            c = c.where(col == val)
    if f.bodegas is not None:
        c = c.where(Existencia.bodega_id.in_(f.bodegas))
    if f.familias is not None:
        c = c.where(Subfamilia.familia_id.in_(f.familias))
    if solo_utilizables:
        c = c.where(Existencia.estado == "DISPONIBLE")
    return list(db.scalars(c.order_by(Existencia.fecha_ingreso, Existencia.id)).unique())


def _fila_existencia(e: Existencia) -> dict:
    lote = e.lote
    return {
        "codigo_publico": e.codigo_publico, "producto_codigo": e.producto.codigo, "producto": e.producto.nombre,
        "familia": e.producto.subfamilia.familia.nombre, "subfamilia": e.producto.subfamilia.nombre,
        "lote": lote.codigo if lote else None, "lote_estado": lote.estado if lote else None,
        "fecha_vencimiento": lote.fecha_vencimiento if lote else None, "bodega": e.bodega.nombre,
        "ubicacion": e.ubicacion.codigo_completo if e.ubicacion else None, "fecha_ingreso": e.fecha_ingreso,
        "dias": (hoy() - dia_local(e.fecha_ingreso)).days, "cantidad": e.cantidad_fisica, "unidad": e.unidad.codigo,
        "costo_unitario": e.costo_unitario, "valor": e.cantidad_fisica * e.costo_unitario,
        "existencia_id": e.id, "producto_id": e.producto_id, "lote_id": e.lote_id, "bodega_id": e.bodega_id,
    }


def _agrupar(filas: list[dict], claves: tuple[str, ...], extra: tuple[str, ...] = ()) -> list[dict]:
    grupos: dict[tuple, dict] = {}
    for f in filas:
        k = tuple(f[c] for c in claves)
        g = grupos.setdefault(k, {**{c: f[c] for c in claves + extra}, "cantidad": Decimal(0), "valor": Decimal(0),
                                  "existencias": 0, "ingreso_mas_antiguo": f["fecha_ingreso"]})
        g["cantidad"] += f["cantidad"]
        g["valor"] += f["valor"]
        g["existencias"] += 1
        g["ingreso_mas_antiguo"] = min(g["ingreso_mas_antiguo"], f["fecha_ingreso"])
    return sorted(grupos.values(), key=lambda g: tuple(str(g[c] or "") for c in claves))


def movimientos(db: Session, f: Filtros, tipos: tuple[str, ...] | None = None, solo_proceso: bool = False) -> list[dict]:
    c = (select(MovimientoInventario).join(Producto, MovimientoInventario.producto_id == Producto.id)
         .join(Subfamilia, Producto.subfamilia_id == Subfamilia.id)
         .join(Existencia, MovimientoInventario.existencia_id == Existencia.id)
         .where(MovimientoInventario.empresa_id == f.empresa_id))
    if tipos:
        c = c.where(MovimientoInventario.tipo_movimiento_id.in_(select(TipoMovimiento.id).where(TipoMovimiento.codigo.in_(tipos))))
    if solo_proceso:
        c = c.where(MovimientoInventario.proceso_id.is_not(None))
    for col, val in ((Subfamilia.familia_id, f.familia_id), (Producto.subfamilia_id, f.subfamilia_id),
                     (MovimientoInventario.producto_id, f.producto_id), (MovimientoInventario.proceso_id, f.proceso_id),
                     (MovimientoInventario.temporada_id, f.temporada_id), (Existencia.bodega_id, f.bodega_id)):
        if val:
            c = c.where(col == val)
    if f.especie or f.productor:
        c = c.join(Proceso, MovimientoInventario.proceso_id == Proceso.id)
        if f.especie:
            c = c.where(Proceso.especie == f.especie)
        if f.productor:
            c = c.where(Proceso.productor == f.productor)
    if f.desde:
        c = c.where(MovimientoInventario.fecha >= inicio_dia_utc(f.desde))
    if f.hasta:
        c = c.where(MovimientoInventario.fecha < inicio_dia_utc(f.hasta + timedelta(days=1)))
    if f.bodegas is not None:
        c = c.where(Existencia.bodega_id.in_(f.bodegas))
    if f.familias is not None:
        c = c.where(Subfamilia.familia_id.in_(f.familias))
    filas = db.scalars(c.order_by(MovimientoInventario.id.desc()).limit(20000)).unique().all()
    usuarios = dict(db.execute(select(Usuario.id, Usuario.nombre)).all())
    procesos = {p.id: p for p in db.scalars(select(Proceso).where(Proceso.id.in_({m.proceso_id for m in filas if m.proceso_id})))} \
        if any(m.proceso_id for m in filas) else {}
    docs = dict(db.execute(select(DocumentoInventario.id, DocumentoInventario.numero).where(
        DocumentoInventario.id.in_({m.documento_id for m in filas if m.documento_id}))).all()) if any(m.documento_id for m in filas) else {}
    anulados = set(db.scalars(select(MovimientoInventario.movimiento_anulado_id).where(
        MovimientoInventario.movimiento_anulado_id.is_not(None), MovimientoInventario.empresa_id == f.empresa_id)))
    from app.modelos import MotivoMovimiento
    rebajas = dict(db.execute(select(MotivoMovimiento.id, MotivoMovimiento.nombre)).all())
    salida = []
    for m in filas:
        e, p = m.existencia, procesos.get(m.proceso_id)
        salida.append({
            "id": m.id, "fecha": m.fecha, "tipo": m.tipo.nombre, "tipo_codigo": m.tipo.codigo,
            "codigo_publico": e.codigo_publico, "producto_codigo": e.producto.codigo, "producto": e.producto.nombre,
            "familia": e.producto.subfamilia.familia.nombre, "subfamilia": e.producto.subfamilia.nombre,
            "lote": e.lote.codigo if e.lote else None, "cantidad": m.cantidad, "unidad": e.unidad.codigo,
            "costo_unitario": m.costo_unitario, "costo_total": m.costo_total, "bodega": e.bodega.nombre,
            "ubicacion": e.ubicacion.codigo_completo if e.ubicacion else None, "proceso": p.codigo if p else None,
            "especie": p.especie if p else None, "productor": p.productor if p else None,
            "documento": docs.get(m.documento_id), "usuario": usuarios.get(m.usuario_id),
            "rebaja": rebajas.get(m.motivo_id), "motivo": m.motivo_texto or m.observacion, "anulado": "SI" if m.id in anulados else "",
            "producto_id": m.producto_id, "proceso_id": m.proceso_id,
        })
    return salida


# ---------------------------------------------------------------- reportes

def r_stock_actual(db, f):
    filas = _agrupar([_fila_existencia(e) for e in existencias_con_stock(db, f)],
                     ("producto_codigo", "producto", "unidad"), ("familia", "subfamilia", "producto_id"))
    reservas = dict(db.execute(select(ReservaStock.producto_id, func.sum(ReservaStock.cantidad - ReservaStock.cantidad_consumida))
                               .where(ReservaStock.empresa_id == f.empresa_id, ReservaStock.estado == "ACTIVA")
                               .group_by(ReservaStock.producto_id)).all())
    minimos = {p.id: p for p in db.scalars(select(Producto).where(Producto.id.in_([g["producto_id"] for g in filas]))).unique()} if filas else {}
    for g in filas:
        res = dec(reservas.get(g["producto_id"]))
        p = minimos[g["producto_id"]]
        g.update(reservado=res, disponible=max(Decimal(0), g["cantidad"] - res), stock_minimo=p.stock_minimo,
                 estado="BAJO MINIMO" if p.stock_minimo is not None and g["cantidad"] <= p.stock_minimo else "")
    return filas


def r_stock_por_bodega(db, f):
    return _agrupar([_fila_existencia(e) for e in existencias_con_stock(db, f)],
                    ("bodega", "producto_codigo", "producto", "unidad"), ("familia", "subfamilia"))


def r_stock_valorizado(db, f):
    return [_fila_existencia(e) for e in existencias_con_stock(db, f)]


def r_stock_por_lote(db, f):
    return _agrupar([_fila_existencia(e) for e in existencias_con_stock(db, f)],
                    ("producto_codigo", "producto", "lote", "bodega", "unidad"), ("lote_estado", "fecha_vencimiento"))


def r_movimientos(db, f):
    return movimientos(db, f)


def r_ingresos(db, f):
    return movimientos(db, f, ("INGRESO", "DEVOLUCION", "AJUSTE_POSITIVO"))


def r_consumos(db, f):
    filas = movimientos(db, f, ("CONSUMO", "SALIDA", "MERMA", "VENCIMIENTO", "AJUSTE_NEGATIVO", "PRESTAMO"))
    for x in filas:
        x["cantidad"], x["costo_total"] = -x["cantidad"], -x["costo_total"]
    return filas


def _analisis(dimension):
    def gen(db, f: Filtros):
        return costeo.analisis(db, f.empresa_id, dimension, desde=f.desde, hasta=f.hasta, temporada_id=f.temporada_id,
                               especie=f.especie, productor=f.productor, proceso_id=f.proceso_id, bodega_id=f.bodega_id,
                               familia_id=f.familia_id, subfamilia_id=f.subfamilia_id, producto_id=f.producto_id,
                               bodegas=f.bodegas, familias=f.familias)
    return gen


def r_consumo_por_proceso(db, f):
    """Detalle por proceso y producto: cantidad neta y costo (las anulaciones ya estan descontadas)."""
    filas = movimientos(db, f, solo_proceso=True)
    grupos: dict[tuple, dict] = {}
    for m in filas:
        g = grupos.setdefault((m["proceso"], m["producto_codigo"]), {
            "proceso": m["proceso"], "especie": m["especie"], "productor": m["productor"],
            "producto_codigo": m["producto_codigo"], "producto": m["producto"], "unidad": m["unidad"],
            "cantidad": Decimal(0), "costo": Decimal(0)})
        g["cantidad"] -= m["cantidad"]
        g["costo"] -= m["costo_total"]
    return sorted((g for g in grupos.values() if g["cantidad"] or g["costo"]), key=lambda g: (g["proceso"], g["producto"]))


def r_costo_por_kilo(db, f):
    c = select(Proceso).where(Proceso.empresa_id == f.empresa_id, Proceso.estado != "ANULADO")
    for col, val in ((Proceso.temporada_id, f.temporada_id), (Proceso.especie, f.especie), (Proceso.productor, f.productor),
                     (Proceso.id, f.proceso_id)):
        if val:
            c = c.where(col == val)
    if f.desde:
        c = c.where(Proceso.fecha >= f.desde)
    if f.hasta:
        c = c.where(Proceso.fecha <= f.hasta)
    salida = []
    for p in db.scalars(c.order_by(Proceso.fecha.desc(), Proceso.id.desc())):
        total, kilos = costeo.costo_proceso(db, p.id), costeo.kilos_base(db, p)
        salida.append({"proceso": p.codigo, "fecha": p.fecha, "estado": p.estado, "especie": p.especie, "variedad": p.variedad,
                       "productor": p.productor, "exportadora": p.exportadora, "kilos": kilos, "estado_kilos": p.estado_kilos,
                       "base_costeo": costeo.ETIQUETAS_BASE.get(costeo.campo_base(db, p), costeo.campo_base(db, p)),
                       "costo": total, "costo_por_kilo": (total / kilos) if kilos else None, "proceso_id": p.id})
    return salida


def r_vencidos(db, f):
    return [x for x in (_fila_existencia(e) for e in existencias_con_stock(db, f))
            if x["fecha_vencimiento"] and x["fecha_vencimiento"] < hoy()]


def r_proximos_a_vencer(db, f):
    limite = hoy() + timedelta(days=f.dias or 30)
    filas = [x for x in (_fila_existencia(e) for e in existencias_con_stock(db, f))
             if x["fecha_vencimiento"] and hoy() <= x["fecha_vencimiento"] <= limite]
    for x in filas:
        x["dias_para_vencer"] = (x["fecha_vencimiento"] - hoy()).days
    return sorted(filas, key=lambda x: x["fecha_vencimiento"])


def rangos_antiguedad(db: Session, empresa_id: int) -> list[tuple[str, int, int | None]]:
    limites = sorted(int(x) for x in (parametro(db, "antiguedad.rangos", empresa_id, [30, 60, 90, 180]) or [30, 60, 90, 180]))
    rangos, desde = [], 0
    for lim in limites:
        rangos.append((f"{desde}-{lim}", desde, lim))
        desde = lim + 1
    rangos.append((f"Mas de {limites[-1]}", desde, None))
    return rangos


def r_antiguedad(db, f):
    rangos = rangos_antiguedad(db, f.empresa_id)
    filas = [_fila_existencia(e) for e in existencias_con_stock(db, f)]
    for x in filas:
        x["rango"] = next(n for n, d, h in rangos if x["dias"] >= d and (h is None or x["dias"] <= h))
    return sorted(filas, key=lambda x: -x["dias"])


def r_prestamos(db, f):
    c = select(Prestamo).where(Prestamo.empresa_id == f.empresa_id)
    if f.producto_id:
        c = c.where(Prestamo.producto_id == f.producto_id)
    if f.desde:
        c = c.where(Prestamo.fecha_prestamo >= inicio_dia_utc(f.desde))
    productos = {p.id: p for p in db.scalars(select(Producto)).unique()}
    salida = []
    for p in db.scalars(c.order_by(Prestamo.id.desc())):
        prod = productos[p.producto_id]
        vencido = p.estado in ("ACTIVO", "PARCIAL") and p.fecha_devolucion_esperada and p.fecha_devolucion_esperada < hoy()
        salida.append({"numero": p.numero, "producto": prod.nombre, "cantidad": p.cantidad, "devuelto": p.cantidad_devuelta,
                       "pendiente": p.cantidad - p.cantidad_devuelta, "unidad": prod.unidad_base.codigo, "persona": p.persona,
                       "area": p.area, "fecha_prestamo": p.fecha_prestamo, "fecha_esperada": p.fecha_devolucion_esperada,
                       "fecha_devolucion": p.fecha_devolucion_real, "estado": "VENCIDO" if vencido else p.estado,
                       "observacion": p.observacion})
    return salida


def r_inventario_fisico(db, f):
    c = select(DetalleConteoInventario, ConteoInventario).join(ConteoInventario, DetalleConteoInventario.conteo_id == ConteoInventario.id) \
        .where(ConteoInventario.empresa_id == f.empresa_id)
    if f.bodega_id:
        c = c.where(ConteoInventario.bodega_id == f.bodega_id)
    if f.desde:
        c = c.where(ConteoInventario.fecha >= inicio_dia_utc(f.desde))
    bodegas = dict(db.execute(select(Bodega.id, Bodega.nombre)).all())
    personas = db.scalars(select(Usuario).where(Usuario.empresa_id == f.empresa_id)).unique().all()
    usuarios = {u.id: u.nombre for u in personas}
    roles = {u.id: ", ".join(r.nombre for r in u.roles) for u in personas}
    salida = []
    for det, conteo in db.execute(c.order_by(ConteoInventario.id.desc(), DetalleConteoInventario.id)).unique():
        e = db.get(Existencia, det.existencia_id)
        salida.append({"conteo": conteo.numero, "fecha": conteo.fecha, "estado": conteo.estado, "bodega": bodegas.get(conteo.bodega_id),
                       "creado_por": usuarios.get(conteo.usuario_id), "contado_por": usuarios.get(det.contado_por),
                       "rol_contador": roles.get(det.contado_por), "autorizado_por": usuarios.get(conteo.autorizado_por),
                       "codigo_publico": e.codigo_publico, "producto": e.producto.nombre, "lote": e.lote.codigo if e.lote else None,
                       "ubicacion": e.ubicacion.codigo_completo if e.ubicacion else None, "cantidad_sistema": det.cantidad_sistema,
                       "cantidad_contada": det.cantidad_contada, "diferencia": det.diferencia, "unidad": e.unidad.codigo,
                       "valor_diferencia": (det.diferencia * e.costo_unitario) if det.diferencia is not None else None,
                       "ajustado": "SI" if det.movimiento_ajuste_id else ""})
    return salida


def r_excepciones_fifo(db, f):
    c = (select(ExcepcionFifo, MovimientoInventario).join(MovimientoInventario, ExcepcionFifo.movimiento_id == MovimientoInventario.id)
         .where(MovimientoInventario.empresa_id == f.empresa_id))
    if f.producto_id:
        c = c.where(MovimientoInventario.producto_id == f.producto_id)
    if f.proceso_id:
        c = c.where(MovimientoInventario.proceso_id == f.proceso_id)
    if f.desde:
        c = c.where(ExcepcionFifo.fecha >= inicio_dia_utc(f.desde))
    if f.hasta:
        c = c.where(ExcepcionFifo.fecha < inicio_dia_utc(f.hasta + timedelta(days=1)))
    usuarios = dict(db.execute(select(Usuario.id, Usuario.nombre)).all())
    salida = []
    for ex, m in db.execute(c.order_by(ExcepcionFifo.id.desc())).unique():
        usada, rec = m.existencia, db.get(Existencia, ex.existencia_recomendada_id)
        if f.bodegas is not None and usada.bodega_id not in f.bodegas:
            continue
        salida.append({"fecha": ex.fecha, "producto": usada.producto.nombre, "existencia_usada": usada.codigo_publico,
                       "lote_usado": usada.lote.codigo if usada.lote else None, "ingreso_usado": usada.fecha_ingreso,
                       "existencia_recomendada": rec.codigo_publico, "lote_recomendado": rec.lote.codigo if rec.lote else None,
                       "ingreso_recomendado": rec.fecha_ingreso, "cantidad": -m.cantidad, "unidad": usada.unidad.codigo,
                       "estrategia": ex.estrategia, "politica": ex.politica, "justificacion": ex.justificacion,
                       "usuario": usuarios.get(ex.usuario_id), "proceso": db.get(Proceso, m.proceso_id).codigo if m.proceso_id else None,
                       "bodega": usada.bodega.nombre})
    return salida


def r_reservas(db, f):
    c = select(ReservaStock).where(ReservaStock.empresa_id == f.empresa_id)
    if f.producto_id:
        c = c.where(ReservaStock.producto_id == f.producto_id)
    if f.proceso_id:
        c = c.where(ReservaStock.proceso_id == f.proceso_id)
    bodegas = dict(db.execute(select(Bodega.id, Bodega.nombre)).all())
    salida = []
    for r in db.scalars(c.order_by(ReservaStock.id.desc())).unique():
        salida.append({"numero": r.numero, "estado": r.estado, "producto": r.producto.nombre, "cantidad": r.cantidad,
                       "consumido": r.cantidad_consumida, "pendiente": (r.cantidad - r.cantidad_consumida) if r.estado == "ACTIVA" else Decimal(0),
                       "unidad": r.producto.unidad_base.codigo, "bodega": bodegas.get(r.bodega_id),
                       "proceso": db.get(Proceso, r.proceso_id).codigo if r.proceso_id else None, "area": r.area,
                       "motivo": r.motivo, "fecha": r.fecha, "fecha_limite": r.fecha_vencimiento})
    return salida


def _valor_por(clave: str, titulo: str):
    def gen(db, f):
        filas = [_fila_existencia(e) for e in existencias_con_stock(db, f)]
        if clave == "proveedor":
            nombres = dict(db.execute(select(Proveedor.id, Proveedor.nombre)).all())
            origen = {e.id: nombres.get(e.proveedor_id, "Sin proveedor") for e in existencias_con_stock(db, f)}
            for x in filas:
                x["proveedor"] = origen[x["existencia_id"]]
        grupos: dict[str, dict] = {}
        for x in filas:
            g = grupos.setdefault(x[clave] or f"Sin {titulo.lower()}", {"clave": x[clave] or f"Sin {titulo.lower()}",
                                                                      "valor": Decimal(0), "existencias": 0})
            g["valor"] += x["valor"]
            g["existencias"] += 1
        total = sum((g["valor"] for g in grupos.values()), Decimal(0))
        for g in grupos.values():
            g["porcentaje"] = (g["valor"] / total * 100) if total else Decimal(0)
        return sorted(grupos.values(), key=lambda g: -g["valor"])
    return gen


C = {
    "prod": [("producto_codigo", "Codigo", "texto"), ("producto", "Producto", "texto")],
    "fam": [("familia", "Familia", "texto"), ("subfamilia", "Subfamilia", "texto")],
    "cant": [("cantidad", "Cantidad", "numero"), ("unidad", "Unidad", "texto")],
}
COLS_ANALISIS = [("clave", "Dimension", "texto"), ("costo", "Costo", "moneda"), ("kilos", "Kilos", "numero"),
                 ("costo_por_kilo", "Costo por kilo", "moneda4"), ("procesos", "Procesos", "numero"), ("estado_kilos", "Kilos", "texto")]
COLS_MOV = [("fecha", "Fecha", "fechahora"), ("tipo", "Tipo", "texto"), ("codigo_publico", "Existencia", "texto"), *C["prod"],
            ("lote", "Lote", "texto"), *C["cant"], ("costo_unitario", "Costo unitario", "moneda4"),
            ("costo_total", "Costo total", "moneda"), ("bodega", "Bodega", "texto"), ("proceso", "Proceso", "texto"),
            ("rebaja", "Tipo de rebaja", "texto"), ("documento", "Documento", "texto"), ("usuario", "Usuario", "texto"),
            ("motivo", "Motivo", "texto"),
            ("anulado", "Anulado", "texto")]
COLS_EXIST = [("codigo_publico", "Existencia", "texto"), *C["prod"], *C["fam"], ("lote", "Lote", "texto"),
              ("fecha_vencimiento", "Vencimiento", "fecha"), ("bodega", "Bodega", "texto"), ("ubicacion", "Ubicacion", "texto"),
              ("fecha_ingreso", "Ingreso", "fecha"), ("dias", "Dias", "numero"), *C["cant"],
              ("costo_unitario", "Costo unitario", "moneda4"), ("valor", "Valor", "moneda")]

REPORTES: dict[str, Reporte] = {r.codigo: r for r in [
    Reporte("stock_actual", "Stock actual", [*C["prod"], *C["fam"], ("cantidad", "Stock fisico", "numero"),
            ("reservado", "Reservado", "numero"), ("disponible", "Disponible", "numero"), ("unidad", "Unidad", "texto"),
            ("stock_minimo", "Minimo", "numero"), ("estado", "Estado", "texto"), ("valor", "Valor", "moneda")], r_stock_actual),
    Reporte("stock_por_bodega", "Stock por bodega", [("bodega", "Bodega", "texto"), *C["prod"], *C["fam"], *C["cant"],
            ("existencias", "Existencias", "numero"), ("valor", "Valor", "moneda")], r_stock_por_bodega),
    Reporte("stock_valorizado", "Stock valorizado por existencia", COLS_EXIST, r_stock_valorizado, requiere_costos=True),
    Reporte("stock_por_lote", "Stock por lote", [*C["prod"], ("lote", "Lote", "texto"), ("lote_estado", "Estado", "texto"),
            ("fecha_vencimiento", "Vencimiento", "fecha"), ("bodega", "Bodega", "texto"), *C["cant"],
            ("ingreso_mas_antiguo", "Ingreso", "fecha"), ("valor", "Valor", "moneda")], r_stock_por_lote),
    Reporte("movimientos", "Movimientos", COLS_MOV, r_movimientos),
    Reporte("ingresos", "Ingresos", COLS_MOV, r_ingresos),
    Reporte("consumos", "Consumos y salidas", COLS_MOV, r_consumos),
    Reporte("consumo_por_proceso", "Consumo por proceso", [("proceso", "Proceso", "texto"), ("especie", "Especie", "texto"),
            ("productor", "Productor", "texto"), *C["prod"], *C["cant"], ("costo", "Costo", "moneda")], r_consumo_por_proceso),
    Reporte("consumo_por_especie", "Consumo por especie", COLS_ANALISIS, _analisis("especie"), requiere_costos=True),
    Reporte("costo_por_kilo", "Costo por kilo por proceso", [("proceso", "Proceso", "texto"), ("fecha", "Fecha", "fecha"),
            ("estado", "Estado", "texto"), ("especie", "Especie", "texto"), ("variedad", "Variedad", "texto"),
            ("productor", "Productor", "texto"), ("kilos", "Kilos", "numero"), ("estado_kilos", "Kilos", "texto"),
            ("base_costeo", "Base", "texto"), ("costo", "Costo", "moneda"), ("costo_por_kilo", "Costo por kilo", "moneda4")],
            r_costo_por_kilo, requiere_costos=True),
    Reporte("costo_por_productor", "Costo por productor", COLS_ANALISIS, _analisis("productor"), requiere_costos=True),
    Reporte("costo_por_variedad", "Costo por variedad", COLS_ANALISIS, _analisis("variedad"), requiere_costos=True),
    Reporte("costo_por_exportadora", "Costo por exportadora", COLS_ANALISIS, _analisis("exportadora"), requiere_costos=True),
    Reporte("costo_por_temporada", "Costo por temporada", COLS_ANALISIS, _analisis("temporada"), requiere_costos=True),
    Reporte("costo_por_semana", "Costo por semana", COLS_ANALISIS, _analisis("semana"), requiere_costos=True),
    Reporte("costo_por_mes", "Costo por mes", COLS_ANALISIS, _analisis("mes"), requiere_costos=True),
    Reporte("costo_por_producto", "Costo por fitosanitario / producto", COLS_ANALISIS, _analisis("producto"), requiere_costos=True),
    Reporte("costo_por_familia", "Costo por familia", COLS_ANALISIS, _analisis("familia"), requiere_costos=True),
    Reporte("costo_por_subfamilia", "Costo por subfamilia", COLS_ANALISIS, _analisis("subfamilia"), requiere_costos=True),
    Reporte("costo_por_bodega", "Costo por bodega", COLS_ANALISIS, _analisis("bodega"), requiere_costos=True),
    Reporte("valor_por_bodega", "Valor de inventario por bodega", [("clave", "Bodega", "texto"), ("valor", "Valor", "moneda"),
            ("porcentaje", "%", "numero"), ("existencias", "Existencias", "numero")], _valor_por("bodega", "Bodega"), requiere_costos=True),
    Reporte("valor_por_familia", "Valor de inventario por familia", [("clave", "Familia", "texto"), ("valor", "Valor", "moneda"),
            ("porcentaje", "%", "numero"), ("existencias", "Existencias", "numero")], _valor_por("familia", "Familia"), requiere_costos=True),
    Reporte("valor_por_subfamilia", "Valor de inventario por subfamilia", [("clave", "Subfamilia", "texto"), ("valor", "Valor", "moneda"),
            ("porcentaje", "%", "numero"), ("existencias", "Existencias", "numero")], _valor_por("subfamilia", "Subfamilia"), requiere_costos=True),
    Reporte("valor_por_proveedor", "Valor de inventario por proveedor", [("clave", "Proveedor", "texto"), ("valor", "Valor", "moneda"),
            ("porcentaje", "%", "numero"), ("existencias", "Existencias", "numero")], _valor_por("proveedor", "Proveedor"), requiere_costos=True),
    Reporte("productos_vencidos", "Productos vencidos", COLS_EXIST, r_vencidos),
    Reporte("proximos_a_vencer", "Proximos a vencer", [*COLS_EXIST[:5], ("fecha_vencimiento", "Vencimiento", "fecha"),
            ("dias_para_vencer", "Dias para vencer", "numero"), *COLS_EXIST[6:8], *C["cant"], ("valor", "Valor", "moneda")],
            r_proximos_a_vencer),
    Reporte("antiguedad", "Antiguedad de stock", [("rango", "Rango (dias)", "texto"), *COLS_EXIST], r_antiguedad),
    Reporte("prestamos", "Prestamos", [("numero", "Numero", "texto"), ("producto", "Producto", "texto"),
            ("cantidad", "Cantidad", "numero"), ("devuelto", "Devuelto", "numero"), ("pendiente", "Pendiente", "numero"),
            ("unidad", "Unidad", "texto"), ("persona", "Persona", "texto"), ("area", "Area", "texto"),
            ("fecha_prestamo", "Prestamo", "fecha"), ("fecha_esperada", "Devolucion esperada", "fecha"),
            ("fecha_devolucion", "Devolucion real", "fecha"), ("estado", "Estado", "texto")], r_prestamos),
    Reporte("inventario_fisico", "Inventario fisico", [("conteo", "Conteo", "texto"), ("fecha", "Fecha", "fecha"),
            ("estado", "Estado", "texto"), ("bodega", "Bodega", "texto"), ("codigo_publico", "Existencia", "texto"),
            ("producto", "Producto", "texto"), ("lote", "Lote", "texto"), ("ubicacion", "Ubicacion", "texto"),
            ("cantidad_sistema", "Sistema", "numero"), ("cantidad_contada", "Contado", "numero"),
            ("diferencia", "Diferencia", "numero"), ("unidad", "Unidad", "texto"), ("valor_diferencia", "Valor diferencia", "moneda"),
            ("ajustado", "Ajustado", "texto"), ("contado_por", "Contado por", "texto"), ("rol_contador", "Rol", "texto"),
            ("creado_por", "Conteo creado por", "texto"), ("autorizado_por", "Autorizado por", "texto")], r_inventario_fisico),
    Reporte("excepciones_fifo", "Excepciones FIFO / FEFO", [("fecha", "Fecha", "fechahora"), ("producto", "Producto", "texto"),
            ("existencia_usada", "Existencia usada", "texto"), ("lote_usado", "Lote usado", "texto"),
            ("ingreso_usado", "Ingreso", "fecha"), ("existencia_recomendada", "Recomendada", "texto"),
            ("lote_recomendado", "Lote recomendado", "texto"), ("ingreso_recomendado", "Ingreso", "fecha"), *C["cant"],
            ("estrategia", "Estrategia", "texto"), ("politica", "Politica", "texto"), ("justificacion", "Justificacion", "texto"),
            ("usuario", "Usuario", "texto"), ("proceso", "Proceso", "texto")], r_excepciones_fifo),
    Reporte("reservas", "Reservas de stock", [("numero", "Numero", "texto"), ("estado", "Estado", "texto"),
            ("producto", "Producto", "texto"), ("cantidad", "Reservado", "numero"), ("consumido", "Consumido", "numero"),
            ("pendiente", "Pendiente", "numero"), ("unidad", "Unidad", "texto"), ("bodega", "Bodega", "texto"),
            ("proceso", "Proceso", "texto"), ("area", "Area", "texto"), ("motivo", "Motivo", "texto"),
            ("fecha", "Fecha", "fecha"), ("fecha_limite", "Fecha limite", "fecha")], r_reservas),
]}
COLUMNAS_COSTO = {"costo_unitario", "costo_total", "valor", "costo", "costo_por_kilo", "valor_diferencia"}


def ejecutar(db: Session, codigo: str, f: Filtros, ver_costos: bool) -> dict:
    rep = REPORTES[codigo]
    columnas = [c for c in rep.columnas if ver_costos or c[0] not in COLUMNAS_COSTO]
    claves = [c[0] for c in columnas]
    filas = rep.generar(db, f)
    ids = ("existencia_id", "producto_id", "lote_id", "proceso_id", "bodega_id", "id")
    return {"codigo": rep.codigo, "titulo": rep.titulo, "generado": json_valor(ahora()),
            "columnas": [{"clave": c[0], "titulo": c[1], "tipo": c[2]} for c in columnas],
            "filas": [{**{k: json_valor(x.get(k)) for k in claves}, **{k: x[k] for k in ids if k in x}} for x in filas],
            "_crudo": [{k: x.get(k) for k in claves} for x in filas]}


# ---------------------------------------------------------------- exportacion

def _texto_celda(valor, tipo: str) -> str:
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%d-%m-%Y %H:%M" if tipo == "fechahora" else "%d-%m-%Y")
    if isinstance(valor, date):
        return valor.strftime("%d-%m-%Y")
    if isinstance(valor, (Decimal, float)):
        v = dec(valor)
        decimales = 0 if tipo == "moneda" else 4 if tipo == "moneda4" else 3
        return f"{v:,.{decimales}f}".replace(",", "X").replace(".", ",").replace("X", ".").rstrip("0").rstrip(",") if decimales else \
            f"{v:,.0f}".replace(",", ".")
    return str(valor)


def a_csv(resultado: dict) -> bytes:
    salida = StringIO()
    w = csv.writer(salida, delimiter=";", lineterminator="\r\n")
    w.writerow([c["titulo"] for c in resultado["columnas"]])
    for fila in resultado["_crudo"]:
        linea = []
        for c in resultado["columnas"]:
            v = fila.get(c["clave"])
            if isinstance(v, (Decimal, float)):
                v = f"{dec(v).normalize():f}".replace(".", ",")
            elif isinstance(v, (date, datetime)):
                v = _texto_celda(v, c["tipo"])
            elif isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
                v = "'" + v  # evita que Excel interprete el texto como formula
            linea.append("" if v is None else v)
        w.writerow(linea)
    return ("﻿" + salida.getvalue()).encode("utf-8")


def a_xlsx(resultado: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = resultado["titulo"][:31]
    ws.append([c["titulo"] for c in resultado["columnas"]])
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F3A5F")
        celda.alignment = Alignment(horizontal="center")
    formatos = {"moneda": "#,##0", "moneda4": "#,##0.00##", "numero": "#,##0.###", "fecha": "DD-MM-YYYY", "fechahora": "DD-MM-YYYY HH:MM"}
    for fila in resultado["_crudo"]:
        valores = []
        for c in resultado["columnas"]:
            v = fila.get(c["clave"])
            if isinstance(v, Decimal):
                v = float(v)
            elif isinstance(v, str) and v[:1] in ("=", "+", "@"):
                v = "'" + v
            valores.append(v)
        ws.append(valores)
    for i, c in enumerate(resultado["columnas"], 1):
        letra = get_column_letter(i)
        ws.column_dimensions[letra].width = max(len(c["titulo"]) + 2, min(40, max(
            (len(_texto_celda(f.get(c["clave"]), c["tipo"])) for f in resultado["_crudo"][:200]), default=8) + 2))
        if c["tipo"] in formatos:
            for celda in ws[letra][1:]:
                celda.number_format = formatos[c["tipo"]]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


def a_pdf(resultado: dict, empresa: str = "") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    salida = BytesIO()
    doc = SimpleDocTemplate(salida, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm, topMargin=10 * mm,
                            bottomMargin=10 * mm, title=resultado["titulo"])
    estilos = getSampleStyleSheet()
    celda = estilos["BodyText"].clone("celda", fontSize=6.5, leading=7.5)
    columnas = resultado["columnas"][:16]  # el ancho de pagina limita las columnas legibles
    datos = [[c["titulo"] for c in columnas]]
    for fila in resultado["_crudo"][:3000]:
        datos.append([Paragraph(_texto_celda(fila.get(c["clave"]), c["tipo"])[:60], celda) for c in columnas])
    tabla = Table(datos, repeatRows=1)
    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, 0), 7), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C8CED6")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F6F8")])]))
    doc.build([Paragraph(f"{resultado['titulo']}", estilos["Title"]),
               Paragraph(f"{empresa} · Generado el {ahora():%d-%m-%Y %H:%M} UTC · {len(resultado['_crudo'])} filas", estilos["Normal"]),
               Spacer(1, 4 * mm), tabla])
    return salida.getvalue()


__all__ = ["REPORTES", "Filtros", "ejecutar", "a_csv", "a_xlsx", "a_pdf", "existencias_con_stock", "movimientos",
           "rangos_antiguedad", "Lote", "Familia", "Ubicacion"]
