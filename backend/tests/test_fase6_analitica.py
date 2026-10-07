"""Fase 6: centro de alertas, dashboard, reportes, exportacion, buscador y trazabilidad."""
from datetime import date, timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.servicios.reportes import REPORTES

HOY = date.today()
VIEJA = (HOY - timedelta(days=100)).isoformat() + "T10:00:00"
RECIENTE = (HOY - timedelta(days=5)).isoformat() + "T10:00:00"


def evaluar(api):
    return api.post("/alertas/evaluar", espera=200)


def abiertas(api, tipo, entidad_id=None):
    return [a for a in api.get("/alertas", tipo=tipo, abiertas=True, tam=500)["items"] if entidad_id is None or a["entidad_id"] == entidad_id]


@pytest.fixture(scope="module")
def escenario(api, base):
    """Stock variado mas un proceso con consumos, para dashboard, reportes y trazabilidad."""
    a = base["producto"]("ANA1", "Fungicida Analitica", stock_minimo=30, stock_critico=10)
    b = base["producto"]("ANA2", "Insecticida Analitica", usa_vencimiento=True)
    ea = base["ingreso"](a["id"], cantidad=2, precio=170000, fecha=VIEJA)
    eb = base["ingreso"](b["id"], cantidad=1, precio=90000, fecha=RECIENTE, fecha_vencimiento=(HOY + timedelta(days=12)).isoformat(),
                         ubicacion="FIT-R01-N02-P01")
    p = api.crear("/procesos", {"codigo": "PROC-ANA", "especie": "Uva", "productor": "Vina Demo", "kilos_procesados": 10000,
                                "fecha": HOY.isoformat()})
    api.post(f"/procesos/{p['id']}/abrir", espera=200)
    c1 = api.post("/consumos", {"proceso_id": p["id"], "existencia_id": ea["existencias"][0]["id"], "cantidad": 5}, espera=201)
    c2 = api.post("/consumos", {"proceso_id": p["id"], "existencia_id": eb["existencias"][0]["id"], "cantidad": 4}, espera=201)
    return dict(a=a, b=b, ea=ea, eb=eb, proceso=p, c1=c1, c2=c2)


# ---------------------------------------------------------------- alertas

def test_niveles_de_stock_generan_y_resuelven_alertas(api, base):
    p = base["producto"]("ALE1", "Producto Alerta", stock_minimo=30, stock_critico=10)
    evaluar(api)
    a = abiertas(api, "STOCK_AGOTADO", p["id"])
    assert len(a) == 1 and a[0]["severidad"] == "ALTA" and a[0]["estado"] == "NUEVA" and a[0]["enlace"] == f"/productos/detalle/?id={p['id']}"
    assert evaluar(api)["creadas"] == 0 and len(abiertas(api, "STOCK_AGOTADO", p["id"])) == 1  # no se duplica mientras siga abierta
    e = base["ingreso"](p["id"], cantidad=5, unidad="L", precio=100, bultos=1)["existencias"][0]
    evaluar(api)
    assert not abiertas(api, "STOCK_AGOTADO", p["id"]) and "5 L" in abiertas(api, "STOCK_CRITICO", p["id"])[0]["descripcion"]
    base["ingreso"](p["id"], cantidad=1)  # 25 L: sobre el critico, bajo el minimo
    evaluar(api)
    assert not abiertas(api, "STOCK_CRITICO", p["id"]) and abiertas(api, "STOCK_BAJO", p["id"])
    base["ingreso"](p["id"], cantidad=1)  # 45 L: normalizado
    r = evaluar(api)
    assert r["resueltas"] >= 1 and not abiertas(api, "STOCK_BAJO", p["id"])
    historial = [x for x in api.get("/alertas", entidad="producto", entidad_id=p["id"])["items"]]
    assert {x["tipo"] for x in historial} == {"STOCK_AGOTADO", "STOCK_CRITICO", "STOCK_BAJO"}
    assert all(x["estado"] == "RESUELTA" and "automatica" in x["observacion"] for x in historial)
    assert e["id"]


def test_vencimientos_y_reglas_configurables(api, base):
    p = base["producto"]("ALE2", "Producto Vence", usa_vencimiento=True)
    en_20 = base["ingreso"](p["id"], fecha_vencimiento=(HOY + timedelta(days=20)).isoformat())["existencias"][0]
    en_50 = base["ingreso"](p["id"], fecha_vencimiento=(HOY + timedelta(days=50)).isoformat())["existencias"][0]
    vencido = base["ingreso"](p["id"], fecha_vencimiento=(HOY - timedelta(days=3)).isoformat())["existencias"][0]
    evaluar(api)
    assert abiertas(api, "PROXIMO_A_VENCER", en_20["lote_id"]) and not abiertas(api, "PROXIMO_A_VENCER", en_50["lote_id"])
    alerta = abiertas(api, "PRODUCTO_VENCIDO", vencido["lote_id"])[0]
    assert alerta["severidad"] == "CRITICA" and api.get(f"/lotes/{vencido['lote_id']}")["estado"] == "VENCIDO"
    # La anticipacion es configurable: con 60 dias tambien avisa el lote que vence en 50.
    regla = next(r for r in api.get("/reglas-alertas") if r["tipo"] == "PROXIMO_A_VENCER")
    api.put(f"/reglas-alertas/{regla['id']}", {"parametros": {"dias": 60}, "severidad": "ALTA"})
    evaluar(api)
    nueva = abiertas(api, "PROXIMO_A_VENCER", en_50["lote_id"])
    assert nueva and nueva[0]["severidad"] == "ALTA"
    # Una regla desactivada deja de generar y cierra sus alertas.
    api.put(f"/reglas-alertas/{regla['id']}", {"activa": False})
    evaluar(api)
    assert not abiertas(api, "PROXIMO_A_VENCER")
    api.put(f"/reglas-alertas/{regla['id']}", {"activa": True, "parametros": {"dias": 30}, "severidad": "MEDIA"})
    api.put(f"/reglas-alertas/{regla['id']}", {"parametros": {"dias": -1}}, espera=400)
    api.put(f"/reglas-alertas/{regla['id']}", {"activa": False}, espera=403, como="digitador")


def test_stock_antiguo_lote_bloqueado_y_prestamo_vencido(api, base):
    p = base["producto"]("ALE3", "Producto Antiguo")
    viejo = base["ingreso"](p["id"], fecha=VIEJA)["existencias"][0]
    nuevo = base["ingreso"](p["id"], fecha=RECIENTE)["existencias"][0]
    api.post(f"/lotes/{nuevo['lote_id']}/estado", {"estado": "CUARENTENA", "motivo": "Analisis"}, espera=200)
    pr = api.crear("/prestamos", {"existencia_id": viejo["id"], "cantidad": 1, "persona": "Juan Soto",
                                  "fecha_devolucion_esperada": (HOY - timedelta(days=2)).isoformat()})
    evaluar(api)
    a = abiertas(api, "STOCK_ANTIGUO", viejo["id"])
    assert a and "100 dias" in a[0]["descripcion"] and a[0]["bodega"] == "Bodega Fitosanitarios"
    assert not abiertas(api, "STOCK_ANTIGUO", nuevo["id"])
    assert abiertas(api, "LOTE_CUARENTENA", nuevo["lote_id"]) and abiertas(api, "PRESTAMO_VENCIDO", pr["id"])
    api.post(f"/prestamos/{pr['id']}/devolver", {}, espera=200)
    api.post(f"/lotes/{nuevo['lote_id']}/estado", {"estado": "DISPONIBLE", "motivo": "Aprobado"}, espera=200)
    evaluar(api)
    assert not abiertas(api, "LOTE_CUARENTENA", nuevo["lote_id"]) and not abiertas(api, "PRESTAMO_VENCIDO", pr["id"])


def test_gestion_de_estados_de_alerta(api):
    a = api.crear("/alertas", {"titulo": "Revisar candado bodega", "descripcion": "Aviso manual", "severidad": "BAJA"})
    assert a["tipo"] == "OTRA" and a["estado"] == "NUEVA"
    assert api.post(f"/alertas/{a['id']}/estado", {"estado": "LEIDA"}, espera=200, como="lector")["estado"] == "LEIDA"
    api.post(f"/alertas/{a['id']}/estado", {"estado": "EN_GESTION"}, espera=403, como="lector")
    assert api.post(f"/alertas/{a['id']}/estado", {"estado": "EN_GESTION"}, espera=200, como="jefe")["estado"] == "EN_GESTION"
    api.post(f"/alertas/{a['id']}/estado", {"estado": "RESUELTA"}, espera=400, como="jefe")  # exige observacion
    r = api.post(f"/alertas/{a['id']}/estado", {"estado": "RESUELTA", "observacion": "Candado cambiado"}, espera=200, como="jefe")
    assert r["estado"] == "RESUELTA" and r["resuelta_por_nombre"] is None and r["fecha_resolucion"] and r["abierta"] is False
    assert api.get("/alertas", estado="RESUELTA")["items"][0]["resuelta_por_nombre"] == "Jefe"
    api.post(f"/alertas/{a['id']}/estado", {"estado": "INVENTADO"}, espera=400)
    b = api.crear("/alertas", {"titulo": "Otra", "severidad": "CRITICA"})
    assert api.post(f"/alertas/{b['id']}/estado", {"estado": "DESCARTADA", "observacion": "Falsa alarma"}, espera=200)["estado"] == "DESCARTADA"
    opciones = api.get("/alertas/opciones")
    assert "STOCK_BAJO" in opciones["tipos"] and opciones["severidades"] == ["INFORMATIVA", "BAJA", "MEDIA", "ALTA", "CRITICA"]
    assert opciones["estados"] == ["NUEVA", "LEIDA", "EN_GESTION", "RESUELTA", "DESCARTADA"]


def test_resumen_de_alertas(api, escenario):
    evaluar(api)
    r = api.get("/alertas/resumen")
    assert r["abiertas"] >= r["nuevas"] >= 1 and r["proximos_a_vencer"] >= 1 and r["vencidos"] >= 1
    assert sum(r["por_tipo"].values()) == r["abiertas"] == sum(r["por_severidad"].values())


def test_canal_de_notificacion_es_extensible(api, base):
    """El centro de alertas admite canales (email, Teams, Telegram, WhatsApp) sin modificar el resto."""
    from app.servicios import alertas as svc
    enviados = []
    svc.registrar_canal("email", lambda a: enviados.append(a.titulo))
    try:
        regla = next(r for r in api.get("/reglas-alertas") if r["tipo"] == "STOCK_AGOTADO")
        api.put(f"/reglas-alertas/{regla['id']}", {"canales": ["EMAIL", "TELEGRAM"]})  # Telegram no registrado: se ignora
        base["producto"]("ALE9", "Producto Notificado", stock_minimo=5)
        evaluar(api)
        assert enviados == ["Sin stock: Producto Notificado"]
        api.put(f"/reglas-alertas/{regla['id']}", {"canales": []})
    finally:
        svc.CANALES.clear()


# ---------------------------------------------------------------- dashboard

def test_dashboard_principal(api, escenario):
    d = api.get("/dashboard")
    existencias = api.get("/existencias", tam=1000)["items"]
    assert d["valor_total_inventario"] == pytest.approx(sum(e["valor_disponible"] for e in existencias))
    assert sum(x["valor"] for x in d["valor_por_bodega"]) == pytest.approx(d["valor_total_inventario"])
    assert d["valor_por_familia"][0]["nombre"] == "Fitosanitarios" and d["valor_por_subfamilia"]
    assert d["stock_total"]["existencias"] == len(existencias)
    assert d["consumo_del_dia"]["costo"] == pytest.approx(5 * 8500 + 4 * 4500) and d["consumo_del_dia"]["movimientos"] == 2
    assert d["consumo_semanal"]["costo"] >= d["consumo_del_dia"]["costo"] and len(d["consumo_periodo"]["serie"]) == 30
    uva = next(x for x in d["costo_por_especie"] if x["clave"] == "Uva")
    assert uva["costo"] == 60500 and uva["costo_por_kilo"] == 6.05
    assert d["costo_por_proceso"][0]["clave"] == "PROC-ANA" and d["costo_por_kilo"]["estado_kilos"] == "PROVISORIO"
    assert {x["nombre"] for x in d["productos_mas_consumidos"]} == {"Fungicida Analitica", "Insecticida Analitica"}
    assert d["ultimos_movimientos"][0]["tipo_codigo"] == "CONSUMO" and len(d["ultimos_movimientos"]) == 10
    assert d["stock_antiguo"]["existencias"] >= 2 and d["stock_antiguo"]["valor_inventario"] > 0
    assert d["proximos_a_vencer"]["cantidad"] >= 1 and d["proximos_a_vencer"]["lotes"][0]["dias"] <= 30
    assert any(p["producto"] == "Fungicida Analitica" for p in d["stock_bajo_minimo"]["productos"]) is False  # 35 L > minimo 30
    ant = d["antiguedad"]
    assert [r["rango"] for r in ant["rangos"]] == ["0-30", "31-60", "61-90", "91-180", "Mas de 180"]
    assert sum(r["existencias"] for r in ant["rangos"]) == len(existencias) and ant["lote_mas_antiguo"]["dias"] >= 100
    assert ant["edad_promedio_dias"] > 0 and d["alertas"]["abiertas"] >= 1


def test_dashboard_filtros_y_rangos_configurables(api, escenario):
    por_producto = api.get("/dashboard", producto_id=escenario["a"]["id"])
    assert por_producto["valor_total_inventario"] == 35 * 8500 and por_producto["stock_total"]["productos"] == 1
    assert api.get("/dashboard", especie="Uva")["costo_por_kilo"]["costo_por_kilo"] == 6.05
    assert api.get("/dashboard", especie="Manzana")["costo_por_proceso"] == []
    assert api.get("/dashboard", bodega_id=escenario["ea"]["bodega_id"])["valor_por_bodega"][0]["nombre"] == "Bodega Fitosanitarios"
    assert api.get("/dashboard", proceso_id=escenario["proceso"]["id"], productor="Vina Demo")["costo_por_proceso"][0]["costo"] == 60500
    ayer = (HOY - timedelta(days=1)).isoformat()
    assert api.get("/dashboard", desde="2020-01-01", hasta=ayer)["consumo_del_dia"]["costo"] == 0
    api.put("/configuraciones/antiguedad.rangos", {"valor": [15, 45]})
    assert [r["rango"] for r in api.get("/dashboard")["antiguedad"]["rangos"]] == ["0-15", "16-45", "Mas de 45"]
    api.put("/configuraciones/antiguedad.rangos", {"valor": [30, 60, 90, 180]})


def test_dashboard_sin_permiso_de_costos(api, escenario):
    d = api.get("/dashboard", como="digitador")
    assert d["valor_total_inventario"] is None and d["valor_por_bodega"] == [] and d["costo_por_kilo"] is None
    assert d["consumo_del_dia"]["costo"] is None and d["stock_antiguo"]["valor_inventario"] is None
    assert d["stock_total"]["existencias"] > 0 and d["antiguedad"]["rangos"][0]["valor_inventario"] is None
    assert all(m["costo_total"] is None for m in d["ultimos_movimientos"])


# ---------------------------------------------------------------- reportes

def test_todos_los_reportes_ejecutan(api, escenario):
    requeridos = {"stock_actual", "stock_por_bodega", "stock_valorizado", "stock_por_lote", "movimientos", "ingresos", "consumos",
                  "consumo_por_proceso", "consumo_por_especie", "costo_por_kilo", "costo_por_productor", "productos_vencidos",
                  "proximos_a_vencer", "prestamos", "inventario_fisico", "excepciones_fifo"}
    disponibles = {r["codigo"] for r in api.get("/reportes")}
    assert requeridos <= disponibles == set(REPORTES)
    for codigo in sorted(disponibles):
        r = api.get(f"/reportes/{codigo}")
        claves = [c["clave"] for c in r["columnas"]]
        assert r["titulo"] and claves and all(set(claves) <= set(f) for f in r["filas"]), codigo
    api.get("/reportes/inexistente", espera=404)


def test_contenido_de_reportes(api, escenario):
    a = escenario["a"]
    fila = next(f for f in api.get("/reportes/stock_actual", producto_id=a["id"])["filas"])
    assert (fila["cantidad"], fila["reservado"], fila["disponible"], fila["valor"], fila["estado"]) == (35, 0, 35, 297500, "")
    consumo = api.get("/reportes/consumos", proceso_id=escenario["proceso"]["id"])["filas"]
    assert sorted(f["costo_total"] for f in consumo) == [18000, 42500] and all(f["cantidad"] > 0 for f in consumo)
    cpp = api.get("/reportes/consumo_por_proceso", proceso_id=escenario["proceso"]["id"])["filas"]
    assert {(f["producto"], f["cantidad"], f["costo"]) for f in cpp} == {("Fungicida Analitica", 5, 42500), ("Insecticida Analitica", 4, 18000)}
    kg = next(f for f in api.get("/reportes/costo_por_kilo")["filas"] if f["proceso"] == "PROC-ANA")
    assert (kg["costo"], kg["kilos"], kg["costo_por_kilo"], kg["estado_kilos"], kg["base_costeo"]) == (60500, 10000, 6.05, "PROVISORIO", "Kilos procesados")
    assert next(f for f in api.get("/reportes/costo_por_productor")["filas"] if f["clave"] == "Vina Demo")["costo"] == 60500
    prox = api.get("/reportes/proximos_a_vencer", dias=15, producto_id=escenario["b"]["id"])["filas"]
    assert len(prox) == 1 and prox[0]["dias_para_vencer"] == 12
    assert api.get("/reportes/proximos_a_vencer", dias=5, producto_id=escenario["b"]["id"])["filas"] == []
    ant = api.get("/reportes/antiguedad", producto_id=a["id"])["filas"]
    assert {f["rango"] for f in ant} == {"91-180"}
    hoy = HOY.isoformat()
    assert len(api.get("/reportes/movimientos", desde=hoy, hasta=hoy, producto_id=a["id"])["filas"]) == 1
    assert len(api.get("/reportes/movimientos", producto_id=a["id"])["filas"]) == 3


def test_exportacion_excel_csv_pdf(api, escenario):
    pid = escenario["a"]["id"]
    xlsx = api.get("/reportes/stock_valorizado", formato="xlsx", producto_id=pid)
    assert "spreadsheetml" in xlsx.headers["content-type"] and "stock_valorizado_" in xlsx.headers["content-disposition"]
    hoja = load_workbook(BytesIO(xlsx.content)).active
    filas = list(hoja.iter_rows(values_only=True))
    assert filas[0][0] == "Existencia" and len(filas) == 3 and sum(f[-1] for f in filas[1:]) == 297500
    assert hoja.freeze_panes == "A2"
    csv = api.get("/reportes/stock_actual", formato="csv", producto_id=pid)
    texto = csv.content.decode("utf-8-sig")
    assert texto.splitlines()[0].startswith("Codigo;Producto;") and "ANA1;Fungicida Analitica" in texto
    pdf = api.get("/reportes/movimientos", formato="pdf", producto_id=pid)
    assert pdf.headers["content-type"] == "application/pdf" and pdf.content[:4] == b"%PDF"
    api.get("/reportes/stock_actual", espera=400, formato="doc")
    assert api.get("/auditoria", entidad="reporte", accion="EXPORTAR")["total"] == 3


def test_reportes_respetan_permisos(api, escenario):
    api.get("/reportes/stock_valorizado", espera=403, como="digitador")  # reporte de costos
    r = api.get("/reportes/stock_actual", como="digitador")
    assert "valor" not in [c["clave"] for c in r["columnas"]] and all("valor" not in f for f in r["filas"])
    assert not any(x["requiere_costos"] for x in api.get("/reportes", como="digitador"))
    api.get("/reportes/stock_actual", espera=403, como="digitador", formato="xlsx")  # ver no implica exportar
    assert api.get("/reportes/stock_actual", formato="csv", como="jefe").content


# ---------------------------------------------------------------- buscador y trazabilidad

def test_buscador_general(api, escenario):
    e = escenario["ea"]["existencias"][0]
    token = api.get(f"/existencias/{e['id']}")["token_qr"]

    def tipos(q, como="admin"):
        return {r["tipo"] for r in api.get("/buscar", q=q, como=como)}

    assert "producto" in tipos("Fungicida Anal") and "producto" in tipos("ana1")
    assert tipos(e["lote"]) == {"lote"} and tipos(e["codigo_publico"]) == {"existencia"}
    assert tipos(f"https://wms/q/{token}") == {"existencia"}
    assert "ubicacion" in tipos("FIT-R01-N02") and "proceso" in tipos("PROC-ANA") and "proceso" in tipos("Vina")
    assert "documento" in tipos(escenario["ea"]["numero"]) and "documento" in tipos("F-1001")
    assert api.get("/buscar", q="x") == [] and all(r["enlace"] for r in api.get("/buscar", q="ana"))


def test_trazabilidad_hacia_adelante_desde_lote(api, escenario):
    e = escenario["ea"]["existencias"][0]
    t = api.get(f"/trazabilidad/lote/{e['lote_id']}")
    assert t["raiz"] == f"lote:{e['lote_id']}"
    assert [c["tipo"] for c in t["columnas"]] == ["documento", "lote", "existencia", "movimiento", "proceso"]
    assert t["resumen"] == {"documento": 1, "lote": 1, "existencia": 2, "movimiento": 3, "proceso": 1}
    enlaces = {(x["origen"], x["destino"]) for x in t["enlaces"]}
    doc, lote, ex, proc = f"documento:{escenario['ea']['id']}", f"lote:{e['lote_id']}", f"existencia:{e['id']}", f"proceso:{escenario['proceso']['id']}"
    consumo = f"movimiento:{escenario['c1']['movimiento']['id']}"
    assert {(doc, lote), (lote, ex), (ex, consumo), (consumo, proc)} <= enlaces  # cadena completa y navegable
    assert [ev["tipo"] for ev in t["linea_tiempo"]] == ["INGRESO", "INGRESO", "CONSUMO"]
    nodos = {n["clave"]: n for c in t["columnas"] for n in c["nodos"]}
    assert nodos[doc]["enlace"] and "FACTURA F-1001" in nodos[doc]["detalle"] and "Agroquimica Demo" in nodos[doc]["detalle"]
    assert nodos[proc]["titulo"] == "PROC-ANA" and nodos[ex]["enlace"] == f"/existencia/?id={e['id']}"


def test_trazabilidad_hacia_atras_desde_proceso(api, escenario):
    t = api.get(f"/trazabilidad/proceso/{escenario['proceso']['id']}")
    assert t["raiz"] == f"proceso:{escenario['proceso']['id']}"
    assert t["resumen"] == {"documento": 2, "lote": 2, "existencia": 2, "movimiento": 2, "proceso": 1}
    assert {(p["producto"], p["cantidad"], p["costo"]) for p in t["productos"]} == {
        ("Fungicida Analitica", 5, 42500), ("Insecticida Analitica", 4, 18000)}
    assert all(p["lotes"] for p in t["productos"])
    sin_costos = api.get(f"/trazabilidad/proceso/{escenario['proceso']['id']}", como="digitador")
    assert all(p["costo"] is None for p in sin_costos["productos"])
    por_existencia = api.get(f"/trazabilidad/existencia/{escenario['eb']['existencias'][0]['id']}")
    assert por_existencia["resumen"]["proceso"] == 1 and por_existencia["resumen"]["documento"] == 1
    assert api.get(f"/trazabilidad/producto/{escenario['a']['id']}")[0]["codigo"] == escenario["ea"]["existencias"][0]["lote"]
    api.get("/trazabilidad/lote/999999", espera=404)
